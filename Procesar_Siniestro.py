import pandas as pd
import re
import os
import getpass
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# Ruta original del archivo
original_txt_file = r'C:\Users\ccortes\UIB COLOMBIA S.A. Corredores de Reaseguros\UIB Seguros - Documentos\Camilo\Nueva carpeta\Siniestros\Listado_Siniestro.txt'

excel_dest_path = r'C:\Users\ccortes\UIB COLOMBIA S.A. Corredores de Reaseguros\UIB Seguros - Documentos\Camilo\Nueva carpeta\Siniestros\SAP Control de pagos y aprobaciones 2025 - Camilo.xlsx'

print(f"🚀 Iniciando procesamiento...")
print(f"📂 Archivo origen: {original_txt_file}")
print(f"📊 Archivo destino: {excel_dest_path}")

# Verificación de existencia de archivos
if not os.path.exists(original_txt_file):
    print(f"❌ ERROR: No se encuentra el archivo de origen: {original_txt_file}")
    input("Presione Enter para salir...")
    exit()

if not os.path.exists(excel_dest_path):
    print(f"❌ ERROR: No se encuentra el archivo Excel de destino: {excel_dest_path}")
    input("Presione Enter para salir...")
    exit()

# 1. Leer todas las líneas, unificar tabuladores
print("📖 Leyendo archivo TXT...")
with open(original_txt_file, 'r', encoding='utf-8') as f:
    raw_lines = f.readlines()

# Reemplazar tabuladores por punto y coma
raw_lines = [line.replace('\t', ';').strip('\n') for line in raw_lines]

# 2. Usar encabezado para determinar el número esperado de columnas
header = raw_lines[0]
header_parts = [col.strip() for col in header.split(';')]
expected_cols = len(header_parts)
clean_header = ';'.join(header_parts)
raw_lines[0] = clean_header

# 3. Unir líneas partidas
fixed_lines = [clean_header]
buffer = ""

for line in raw_lines[1:]:
    if buffer:
        buffer += " " + line
    else:
        buffer = line

    current_cols = len(buffer.split(';'))

    if current_cols == expected_cols:
        fixed_lines.append(buffer)
        buffer = ""
    elif current_cols > expected_cols:
        print(f"⚠️ Línea demasiado larga: {buffer[:100]}...")
        fixed_lines.append(buffer)
        buffer = ""

# Manejo de buffer remanente
if buffer:
    fixed_lines.append(buffer)

# 4. Guardar archivo intermedio (respaldo limpio)
cleaned_txt_file = original_txt_file.replace('.txt', '_cleaned_backup.txt')
with open(cleaned_txt_file, 'w', encoding='utf-8') as f:
    f.write('\n'.join(fixed_lines))

# 5. Leer con pandas
df = pd.read_csv(cleaned_txt_file, delimiter=';', encoding='utf-8', on_bad_lines='skip', low_memory=False)

# 6. Corregir encabezado mal alineado (columnas 48 a 69)
if len(df.columns) > 48 and df.columns[48].startswith('Unnamed'):
    print("⚠️ Encabezado mal alineado entre columnas 48 a 69. Realineando...")
    new_columns = list(df.columns)
    for i in range(48, 69):
        new_columns[i] = df.columns[i + 1]
    df.columns = new_columns
    df.drop(df.columns[69], axis=1, inplace=True)

# Limpiar caracteres ilegales para Excel (caracteres de control)
def remove_illegal_chars(val):
    if isinstance(val, str):
        # Elimina caracteres de control que rompen openpyxl/Excel
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', val)
    return val

print("🧹 Limpiando caracteres especiales para Excel...")
df = df.map(remove_illegal_chars)

# 7. Unificar columnas ETAPAS y Observaciones
if df.shape[1] >= 55:
    # Vectorización para mayor velocidad y evitar errores de KeyError por posición
    etapas_1 = df.iloc[:, 51].astype(str).replace('nan', '').str.strip()
    etapas_2 = df.iloc[:, 53].astype(str).replace('nan', '').str.strip()
    df.iloc[:, 51] = (etapas_1 + " " + etapas_2).str.strip()

    obs_1 = df.iloc[:, 52].astype(str).replace('nan', '').str.strip()
    obs_2 = df.iloc[:, 54].astype(str).replace('nan', '').str.strip()
    df.iloc[:, 52] = (obs_1 + " " + obs_2).str.strip()

    # Eliminar las columnas sobrantes (53 y 54)
    # Usamos indices de columnas actuales para evitar errores si los nombres cambiaron
    cols_to_drop = [df.columns[53], df.columns[54]]
    df.drop(columns=cols_to_drop, inplace=True)

# 8. Guardar archivo final .txt y .xlsx
final_txt_path = original_txt_file.replace('.txt', '_final.txt')
xlsx_file_path = original_txt_file.replace('.txt', '.xlsx')

df.to_csv(final_txt_path, sep=';', index=False, encoding='utf-8')
df.to_excel(xlsx_file_path, index=False, engine='openpyxl')

print(f'\n✅ TXT corregido guardado como:\n{final_txt_path}')
print(f'✅ Excel generado correctamente:\n{xlsx_file_path}')

# 9. Insertar en archivo destino (BASE SAP desde B27)
print(f"📥 Insertando datos en el Excel de destino: {excel_dest_path}")
wb = load_workbook(excel_dest_path)
if 'BASE SAP' not in wb.sheetnames:
    print(f"⚠️ La hoja 'BASE SAP' no existe. Hojas disponibles: {wb.sheetnames}")
    input("Presione Enter para salir...")
    exit()

ws = wb['BASE SAP']
start_row = 27
start_col = 2  # B

print(f"📝 Escribiendo {len(df)} filas comenzando en B27...")
for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=False), start=start_row):
    for c_idx, value in enumerate(row, start=start_col):
        ws.cell(row=r_idx, column=c_idx, value=value)

print("💾 Guardando cambios en el archivo Excel...")
wb.save(excel_dest_path)
print("✅ Datos pegados exitosamente desde B27 en BASE SAP.")

# 10. Registrar log de actualización
log_path = os.path.join(os.path.dirname(original_txt_file), "log_actualizacion.txt")
now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
user = getpass.getuser()

try:
    with open(log_path, 'w', encoding='utf-8') as log_f:
        log_f.write(f"Actualizado el: {now} por el usuario: {user}")
    print(f"✅ Log actualizado en: {log_path}")
except Exception as e:
    print(f"⚠️ No se pudo escribir el log: {e}")

input("\nPresione Enter para salir...")
